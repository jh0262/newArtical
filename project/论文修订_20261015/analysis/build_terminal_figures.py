"""Extract original DOCX tables, then draw terminal statistics only. No optimizer calls."""
from pathlib import Path
from datetime import datetime
import csv, hashlib, json, math, re, html
from docx import Document
from lxml import etree
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import pypdfium2 as pdfium

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / 'PADE_Mechanical_Optimization_Full_Paper-v2.docx'
FIG = ROOT / 'figures'
FIG.mkdir(exist_ok=True)
AN = ROOT / 'analysis'
before = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
pdfmetrics.registerFont(TTFont('Arial', r'C:\Windows\Fonts\arial.ttf'))
pdfmetrics.registerFont(TTFont('Arial-Bold', r'C:\Windows\Fonts\arialbd.ttf'))
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

def flat(e):
    tag = etree.QName(e).localname
    if e.tag in (f'{{{W}}}t', f'{{{M}}}t'):
        return e.text or ''
    if e.tag == f'{{{M}}}sSup':
        return flat(e.find(f'{{{M}}}e')) + '^(' + flat(e.find(f'{{{M}}}sup')) + ')'
    if e.tag == f'{{{M}}}sSub':
        return flat(e.find(f'{{{M}}}e')) + '_(' + flat(e.find(f'{{{M}}}sub')) + ')'
    return ''.join(flat(c) for c in e)

def num(s):
    s = s.replace('−', '-').replace(' ', '').replace('×', '*').replace('^(', '**(')
    if re.fullmatch(r'[+\-0-9.eE*()]+', s):
        return float(eval(s, {'__builtins__': {}}, {}))
    raise ValueError(f'Not a numeric table cell: {s!r}')

doc = Document(SOURCE)
tables = []
for ti, table in enumerate(doc.tables, 1):
    rows = [[flat(c._tc).strip() for c in row.cells] for row in table.rows]
    tables.append(rows)
    with (AN / f'source_table_{ti}.csv').open('w', encoding='utf-8-sig', newline='') as f:
        csv.writer(f).writerows(rows)
    # Exact source cell XML also preserves every OMML formula and superscript.
    (AN / f'source_table_{ti}.xml').write_bytes(etree.tostring(table._tbl, encoding='UTF-8', xml_declaration=True))

payload = {'source': str(SOURCE), 'source_sha256': before,
           'extraction': 'python-docx; w:t and OMML m:t in document order; superscript as ^(exponent)',
           'tables': [{'table': i+1, 'rows': rows, 'row_count': len(rows), 'column_count': len(rows[0])}
                      for i, rows in enumerate(tables)]}
(AN / 'source_tables.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
assert len(tables) == 5
assert [len(t) for t in tables] == [24, 10, 8, 8, 10]

INK = '#202632'; GRID = '#dce1e8'; BLUE = '#376a94'; RED = '#bc3434'; MUTED = '#586475'

class Drawing:
    def __init__(self, name, width, height):
        self.name, self.w, self.h = name, width, height
        self.pdf = canvas.Canvas(str(FIG / f'{name}.pdf'), pagesize=(width, height))
        self.pdf.setTitle(name + ' - source-table terminal statistics')
        self.parts = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">', '<rect width="100%" height="100%" fill="white"/>']
    def line(self, x1, y1, x2, y2, color=INK, width=.8):
        self.pdf.setStrokeColor(color); self.pdf.setLineWidth(width); self.pdf.line(x1,y1,x2,y2)
        self.parts.append(f'<line x1="{x1}" y1="{self.h-y1}" x2="{x2}" y2="{self.h-y2}" stroke="{color}" stroke-width="{width}"/>')
    def text(self, x,y,t,size=11,color=INK,bold=False,anchor='start'):
        font = 'Arial-Bold' if bold else 'Arial'
        self.pdf.setFont(font,size); self.pdf.setFillColor(color)
        method = {'start':self.pdf.drawString,'middle':self.pdf.drawCentredString,'end':self.pdf.drawRightString}[anchor]
        method(x,y,str(t))
        self.parts.append(f'<text x="{x}" y="{self.h-y}" font-family="Arial, sans-serif" font-size="{size}" font-weight="{700 if bold else 400}" fill="{color}" text-anchor="{anchor}">{html.escape(str(t))}</text>')
    def dot(self,x,y,color,r=4):
        self.pdf.setFillColor(color); self.pdf.setStrokeColor(color); self.pdf.circle(x,y,r,fill=1,stroke=0)
        self.parts.append(f'<circle cx="{x}" cy="{self.h-y}" r="{r}" fill="{color}"/>')
    def finish(self):
        self.pdf.showPage(); self.pdf.save()
        (FIG/f'{self.name}.svg').write_text('\n'.join(self.parts+['</svg>']), encoding='utf-8')
        p = pdfium.PdfDocument(str(FIG/f'{self.name}.pdf'))
        p[0].render(scale=2.5).to_pil().save(FIG/f'{self.name}.png')
        p.close()

selected = [(1,'Sphere'),(5,'Rosenbrock'),(8,'Schwefel 2.26'),(9,'Rastrigin'),(11,'Griewank'),(23,'Shekel 10')]
plotted2=[]
d=Drawing('figure2',1080,660)
for k,(fi,name) in enumerate(selected):
    row=tables[0][fi]
    vals=[num(row[4]),num(row[7])]
    raw=[row[4],row[7]]
    positive=all(v>0 for v in vals)
    bx=12+(k%3)*355; by=360 if k<3 else 68
    d.text(bx+18,by+254,f'({chr(97+k)}) F{fi}: {name}',23,bold=True)
    x0,y0,pw,ph=bx+94,by+36,232,174
    if positive:
        lo=math.floor(math.log10(min(vals)))-1; hi=math.ceil(math.log10(max(vals)))+1
        def yy(v): return y0+(math.log10(v)-lo)/(hi-lo)*ph
        step=max(1, math.ceil((hi-lo)/5))
        ticks=list(range(lo,hi+1,step))
        if hi not in ticks: ticks.append(hi)
        for t in ticks:
            y=y0+(t-lo)/(hi-lo)*ph
            d.line(x0,y,x0+pw,y,GRID,.8); d.text(x0-10,y-7,f'1e{t}',20,anchor='end',color=MUTED)
        scale='Mean (log scale)'
    else:
        span=max(vals)-min(vals); lo=min(vals)-span*.48; hi=max(vals)+span*.48
        def yy(v): return y0+(v-lo)/(hi-lo)*ph
        for i in range(5):
            v=lo+(hi-lo)*i/4; y=yy(v)
            d.line(x0,y,x0+pw,y,GRID,.6)
            d.text(x0-10,y-7,f'{v:.0f}' if fi==8 else f'{v:.2f}',20,anchor='end',color=MUTED)
        scale='Mean (linear scale)'
    d.line(x0,y0,x0,y0+ph); d.line(x0,y0,x0+pw,y0)
    d.text(x0,y0+ph+12,scale,20,color=MUTED)
    for j,(value,color,label) in enumerate(zip(vals,[RED,BLUE],['PADE','Standard DE'])):
        x=x0+54+j*121; y=yy(value)
        d.dot(x,y,color,6)
        textval=f'{value:.5g}'
        d.text(x,y+15,textval,22,color,bold=True,anchor='middle')
        d.text(x,y0-30,label,20,anchor='middle')
    plotted2.append({'function':fi,'source_table':1,'source_row':fi+1,'metric':'Mean','PADE':vals[0],'Standard_DE':vals[1],'raw':raw,'scale':'log' if positive else 'linear'})
d.finish()

labels3=['Coello (GA)','Deb (GA)','Ray & Liew','CPSO','HPSO','SCA','PADE']
labels4=['CONLIN','MMA','CS','SCA','m-SCA','SCAHGM','PADE']
labels5=['Coello (GA)','Coello & Montes','CPSO','HPSO','SCAHGM','PADE']
spec=[('Welded beam',3,tables[2][1:],5,None,labels3,'Best cost'),
      ('Cantilever beam',4,tables[3][1:],6,None,labels4,'Best objective'),
      ('Pressure vessel',5,tables[4][4:],2,4,labels5,'Mean cost +/- Std')]
d=Drawing('figure3',1230,660)
d.text(615,626,'Engineering design benchmark results',28,bold=True,anchor='middle')
plotted3=[]
for k,(title,tableid,rows,col,errcol,labels,metric) in enumerate(spec):
    bx=8+k*405
    d.text(bx+12,576,f'({chr(97+k)}) {title}',26,bold=True)
    vals=[num(r[col]) for r in rows]; errs=[num(r[errcol]) if errcol else 0 for r in rows]
    low=min(v-e for v,e in zip(vals,errs)); high=max(v+e for v,e in zip(vals,errs)); span=high-low
    low-=span*.18; high+=span*.18
    x0,y0,pw,ph=bx+184,130,197,400
    xx=lambda v: x0+(v-low)/(high-low)*pw
    tick_count=3 if tableid==5 else 2
    for i in range(tick_count):
        value=low+(high-low)*i/(tick_count-1); x=xx(value)
        d.line(x,y0-9,x,y0+ph+7,GRID,.6)
        if tableid==3: lab=f'{value:.3f}'
        elif tableid==4: lab=f'{value:.6f}'
        else: lab=f'{value:.0f}'
        d.text(x,y0-36,lab,23,anchor='middle',color=MUTED)
    for i,(r,label,v,e) in enumerate(zip(rows,labels,vals,errs)):
        y=y0+ph-i*ph/(len(rows)-1); color=RED if label=='PADE' else BLUE
        d.text(x0-12,y-7,label,22,anchor='end',bold=label=='PADE',color=color)
        if e:
            a,b=xx(v-e),xx(v+e); d.line(a,y,b,y,color,2); d.line(a,y-6,a,y+6,color,1.8); d.line(b,y-6,b,y+6,color,1.8)
        d.dot(xx(v),y,color,6)
        rawval=r[col]
        d.text(x0+pw,y+18,rawval,24,color,anchor='end')
        plotted3.append({'problem':title,'algorithm':label,'source_table':tableid,'metric':'Mean' if errcol else 'Best','value':v,'value_raw':rawval,'Std':e if errcol else None,'Std_raw':r[errcol] if errcol else None})
    d.line(x0,y0-9,x0+pw,y0-9)
    d.text(x0+pw/2,y0-79,metric,24,anchor='middle')
d.finish()

assert before == hashlib.sha256(SOURCE.read_bytes()).hexdigest(), 'Source modified'
audit={'source_sha256':before,'source_unchanged':True,'figure2':plotted2,'figure3':plotted3,
       'rendering':'ReportLab vector PDF/SVG; pypdfium2 PNG at 180 dpi (2.5 pixels/point)',
       'word_embedding':{'width_cm':14.8,'figure2_height_cm':14.8*660/1080,'figure3_height_cm':14.8*660/1230,'figure2_min_font_pt':20*14.8/2.54*72/1080,'figure3_min_font_pt':22*14.8/2.54*72/1230},
       'not_performed':['optimizer execution','new simulation','trajectory fitting','new significance tests'],
       'files':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(FIG.iterdir()) if p.name.startswith(('figure2.','figure3.'))}}
(AN/'terminal_figures_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'tables':len(tables),'rows':[len(t) for t in tables],'figure2_points':len(plotted2)*2,'figure3_points':len(plotted3),'source_unchanged':True,'outputs':[str(p) for p in sorted(FIG.iterdir())]},ensure_ascii=False,indent=2))
