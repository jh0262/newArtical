"""Editable manuscript with native Word equations, created by bundled runtime."""
from pathlib import Path
from copy import deepcopy
import re,json,importlib.util
from lxml import etree as E
from docx import Document
from docx.shared import Inches,Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent
spec=importlib.util.spec_from_file_location('oldbuilder',ROOT.parent/'manuscript_step5/build_docx.py');o=importlib.util.module_from_spec(spec);spec.loader.exec_module(o)
m,mi,mn,mo,row,sub,sup,frac,br=o.m,o.mi,o.mn,o.mo,o.row,o.sub,o.sup,o.frac,o.bracket
def norm(x):return row(mo('‖'),x,mo('‖'))
def summ(low='l=1',high='K'):return m('munderover',mo('∑'),mi(low),mi(high))
def call(name,x):return row(mi(name),br(x))
w=sub(mi('w'),'l');W=lambda n:sub(mi('W'),str(n));b=lambda n:sub(mi('b'),str(n))
eqs=[
row(o.dot('q'),mo('='),mi('v'),mo(','),o.dot('v'),mo('='),sub(mi('f'),'0'),br(mi('q,v,u;θ')),mo('+'),sup(mi('M'),'-1'),br(mi('q;θ')),sub(mi('r'),'ψ'),br(mi('z'))),
row(call('f',mi('z')),mo('='),b(3),mo('+'),W(3),call('φ',row(b(2),mo('+'),W(2),call('φ',row(b(1),mo('+'),W(1),mi('z'))))),mo('+'),sup(mi('a'),'T'),mi('z')),
row(summ(),sup(norm(w),2),mo('='),frac(sup(norm(mi('w')),2),mi('K')),mo('+'),summ(),sup(norm(row(w,mo('−'),frac(mi('w'),mi('K')))),2)),
row(sup(mi('w'),'+'),mo('='),mi('w'),mo('−'),mi('η'),br(row(mi('Kg'),mo('+'),mi('λw')))),
row(sub(mi('v'),'j'),br(row(mi('t'),mo('−'),mi('τ'))),mo('−'),sub(mi('v'),'j'),br(mi('t')),mo('='),mo('−'),m('msubsup',mo('∫'),row(mi('t'),mo('−'),mi('τ')),mi('t')),sub(mi('a'),'j'),br(mi('s')),mi('ds')),
row(m('mover',mi('Γ'),mo('^')),mo('='),sup(mi('B'),'†'),sub(mi('Z'),'past'),mo(','),mi('R'),mo('='),sub(mi('Z'),'past'),mo('−'),mi('B'),m('mover',mi('Γ'),mo('^')),mo(','),mi('ρ'),mo('='),frac(m('msubsup',norm(mi('R')),mi('F'),mn(2)),m('msubsup',norm(row(sub(mi('Z'),'past'),mo('−'),sub(m('mover',mi('Z'),mo('¯')),'past,train'))),mi('F'),mn(2)))),
row(sub(mi('L'),'a'),mo('='),frac(mn(1),mn(6)),sup(br(row(mi('r'),mo('−'),sup(mi('r'),'*'))),'T'),sup(mi('M'),'-T'),sup(sub(mi('S'),'a'),'-2'),sup(mi('M'),'-1'),br(row(mi('r'),mo('−'),sup(mi('r'),'*')))),
row(sub(mi('e'),'k+1'),mo('≃'),sub(mi('A'),'k'),sub(mi('e'),'k'),mo('+'),sub(mi('B'),'k'),sub(mi('δr'),'k'),mo(','),sub(mi('e'),'q,h'),mo('≃'),sub(mi('C'),'q'),summ('i=0','h−1'),call('Φ',mi('h,i+1')),sub(mi('B'),'i'),sub(mi('δr'),'i')),
o.eqs[3]
]
transform=E.XSLT(E.parse(r'C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL'))
def omml(content):
 # An unmarked mover is converted to a limit, not an accent: dots can disappear.
 content=content.replace('<mover>','<mover accent="true">')
 return deepcopy(transform(E.fromstring(('<math xmlns="http://www.w3.org/1998/Math/MathML">'+content+'</math>').encode())).getroot())
inline={'rψ':sub(mi('r'),'ψ'),'f₀':sub(mi('f'),'0'),'S_a':sub(mi('S'),'a'),'S_q':sub(mi('S'),'q'),'s_q,j':sub(mi('s'),'q,j'),'Z_past':sub(mi('Z'),'past'),'e_(o,h,j)':sub(mi('e'),'o,h,j'),'I_(G,j)':sub(mi('I'),'G,j'),'J_G':sub(mi('J'),'G'),'1_K':sub(mn(1),'K'),'w_l':sub(mi('w'),'l')}
def rich(p,text):
 for t in re.split('('+'|'.join(re.escape(k) for k in inline)+')',text):
  if t in inline:p._p.append(omml(inline[t]))
  elif t:p.add_run(t)
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.3)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体' if name in ['Normal','Caption'] else '黑体')
st=doc.styles['Normal'];st.font.size=Pt(10.5);st.paragraph_format.line_spacing=1.2;st.paragraph_format.space_after=Pt(5);st.paragraph_format.widow_control=True
for name,size in [('Title',17),('Heading 1',13),('Heading 2',11.5)]:
 st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(9.5);doc.styles['Caption'].font.italic=False
for e in doc.styles.element.xpath('.//w:pBdr'):e.getparent().remove(e)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
doc.core_properties.title='重载机器人短历史动力学补偿的增益归因分析';doc.core_properties.author='';doc.core_properties.last_modified_by=''
blocks=re.split(r'\n\s*\n',(ROOT/'论文修订稿.md').read_text(encoding='utf8').strip());eqi=0;figures=0
for block in blocks:
 if block.startswith('$$'):
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p._p.append(omml(eqs[eqi]));p.add_run('    ('+str(eqi+1)+')');eqi+=1;continue
 if block.startswith('|'):
  rs=[r.strip().strip('|').split('|') for r in block.splitlines()];rs=[r for r in rs if not all(re.fullmatch(r'\s*:?-+:?\s*',x) for x in r)]
  table=doc.add_table(rows=0,cols=len(rs[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
  widths={4:[3.0,4.5,2.0,7.0],6:[2.0,1.6,2.7,2.4,4.2,3.6],5:[2.3,3.4,3.4,3.4,3.0]}.get(len(rs[0]),[16.5/len(rs[0])]*len(rs[0]))
  if 'NRMSE' in ''.join(rs[0]) and len(rs[0])==4:widths=[3.4,4.35,4.35,4.4]
  for i,r in enumerate(rs):
   cells=table.add_row().cells
   for c,txt,width in zip(cells,r,widths):
    c.width=Cm(width);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr()
    marg=OxmlElement('w:tcMar')
    for edge in ['top','left','bottom','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');marg.append(e)
    pr.append(marg);border=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');border.append(e)
    pr.append(border)
    if i==0:
     sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAEAEA');pr.append(sh)
    p=c.paragraphs[0];p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
    if '积分步长' in ''.join(rs[0]) and i<len(rs)-1:p.paragraph_format.keep_with_next=True
    run=p.add_run(txt.strip());run.font.size=Pt(8.5);run.bold=i==0
   trpr=table.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
   if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for c,width in zip(table.columns,widths):c.width=Cm(width)
  doc.add_paragraph().paragraph_format.space_after=Pt(0);continue
 if block.startswith('!['):
  match=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',block);p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
  run=p.add_run();run.add_picture(str(ROOT/match[2]),width=Cm(16.3));run._r.xpath('.//wp:docPr')[0].set('descr',match[1])
  cap=doc.add_paragraph(match[1],style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;figures+=1;continue
 if block.startswith('# '):
  p=doc.add_paragraph(block[2:],style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 elif block.startswith('### '):doc.add_paragraph(block[4:],style='Heading 2')
 elif block.startswith('## '):doc.add_paragraph(block[3:],style='Heading 1')
 elif re.match(r'^表\d ',block):
  p=doc.add_paragraph(block,style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
 else:
  p=doc.add_paragraph();rich(p,block)
  if block.startswith('Gain attribution'):
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   for r in p.runs:r.font.size=Pt(11);r.bold=True
  elif re.match(r'^\[\d+\]',block):
   p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(3)
   for r in p.runs:r.font.size=Pt(9)
  elif not re.match('^(关键词|Keywords)',block):p.paragraph_format.first_line_indent=Pt(21)
assert eqi==9
out=ROOT/'重载机器人短历史补偿_论文修订稿.docx';doc.save(out)
print(json.dumps(dict(path=str(out),equations=eqi,figures=figures,tables=len(doc.tables),paragraphs=len(doc.paragraphs)),ensure_ascii=False))
